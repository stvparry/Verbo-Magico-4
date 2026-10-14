from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

NAVY = "17365D"
BLUE = "2F75B5"
GREEN = "538135"
PURPLE = "7030A0"
ORANGE = "C65911"
TEAL = "008C95"
LIGHT = "D9EAF7"
PALE = "EEF5FB"
GREY = "E7E6E6"
WHITE = "FFFFFF"

PRONOUNS = [
    "yo", "tú", "vos", "él", "ella", "usted",
    "nosotros", "nosotras", "vosotros", "vosotras",
    "ellos", "ellas", "ustedes",
]

DISPLAY_PERSONS = [
    ("yo", ["yo"]),
    ("tú", ["tú"]),
    ("él / ella / usted", ["él", "ella", "usted"]),
    ("nosotros / nosotras", ["nosotros", "nosotras"]),
    ("vosotros / vosotras", ["vosotros", "vosotras"]),
    ("ellos / ellas / ustedes", ["ellos", "ellas", "ustedes"]),
]

CORE_TENSES = [
    ("indicativo", "presente", "Presente"),
    ("indicativo", "pretérito-perfecto-simple", "Pretérito perfecto simple"),
    ("indicativo", "pretérito-imperfecto", "Pretérito imperfecto"),
    ("indicativo", "futuro", "Futuro simple"),
    ("condicional", "presente", "Condicional simple"),
]

FULL_SECTIONS = [
    (
        "INDICATIVE • SIMPLE TENSES",
        BLUE,
        [
            ("indicativo", "presente", "Presente"),
            ("indicativo", "pretérito-perfecto-simple", "Pretérito perfecto simple"),
            ("indicativo", "pretérito-imperfecto", "Pretérito imperfecto"),
            ("indicativo", "futuro", "Futuro simple"),
            ("condicional", "presente", "Condicional simple"),
        ],
    ),
    (
        "INDICATIVE • COMPOUND TENSES",
        GREEN,
        [
            ("indicativo", "pretérito-perfecto-compuesto", "Pretérito perfecto compuesto"),
            ("indicativo", "pretérito-pluscuamperfecto", "Pretérito pluscuamperfecto"),
            ("indicativo", "futuro-perfecto", "Futuro perfecto"),
            ("condicional", "perfecto", "Condicional perfecto"),
        ],
    ),
    (
        "SUBJUNCTIVE",
        PURPLE,
        [
            ("subjuntivo", "presente", "Presente de subjuntivo"),
            ("subjuntivo", "pretérito-imperfecto-1", "Imperfecto de subjuntivo (-ra)"),
            ("subjuntivo", "pretérito-imperfecto-2", "Imperfecto de subjuntivo (-se)"),
            ("subjuntivo", "pretérito-perfecto", "Pretérito perfecto de subjuntivo"),
            ("subjuntivo", "pretérito-pluscuamperfecto-1", "Pluscuamperfecto de subjuntivo (-ra)"),
            ("subjuntivo", "pretérito-pluscuamperfecto-2", "Pluscuamperfecto de subjuntivo (-se)"),
        ],
    ),
]


def normalise(payload: dict) -> dict:
    if not isinstance(payload, dict):
        raise ValueError("The conjugation response was not a JSON object.")

    if isinstance(payload.get("value"), dict) and "moods" in payload["value"]:
        return payload["value"]

    if "moods" in payload:
        return payload

    for key in ("data", "result", "conjugation"):
        value = payload.get(key)
        if isinstance(value, dict) and "moods" in value:
            return value

    raise ValueError(f"Unexpected response keys: {list(payload.keys())}")


def strip_pronoun(text: str) -> str:
    text = str(text).strip()
    for pronoun in sorted(PRONOUNS, key=len, reverse=True):
        prefix = pronoun + " "
        if text.casefold().startswith(prefix.casefold()):
            return text[len(prefix):]
    return text


def entries_for(data: dict, mood: str, tense: str) -> list:
    try:
        return data["moods"][mood][tense]
    except KeyError as exc:
        raise ValueError(f"Missing conjugation data: {mood} / {tense}") from exc


def forms_for(data: dict, mood: str, tense: str) -> list[str]:
    by_pronoun = {}
    for entry in entries_for(data, mood, tense):
        pronoun = entry.get("pr")
        forms = entry.get("c") or []
        if pronoun and forms and pronoun not in by_pronoun:
            by_pronoun[pronoun] = strip_pronoun(forms[0])

    output = []
    for _, candidates in DISPLAY_PERSONS:
        values = []
        for pronoun in candidates:
            value = by_pronoun.get(pronoun)
            if value and value not in values:
                values.append(value)
        output.append(" / ".join(values))
    return output


def nonfinite(data: dict, mood: str, tense: str) -> str:
    entries = entries_for(data, mood, tense)
    forms = entries[0].get("c") or []
    if not forms:
        return ""
    return strip_pronoun(forms[0])


def fill(cell, colour: str) -> None:
    cell.fill = PatternFill("solid", fgColor=colour)


def style_title(ws, text: str, end_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    cell = ws.cell(1, 1, text)
    cell.font = Font(name="Aptos Display", size=20, bold=True, color=WHITE)
    fill(cell, NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34
    ws.sheet_view.showGridLines = False


def section_bar(ws, row: int, text: str, colour: str, end_col: int = 7) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_col)
    cell = ws.cell(row, 1, text)
    cell.font = Font(bold=True, size=13, color=WHITE)
    fill(cell, colour)


def build_quick_reference(wb: Workbook, verb: str, data: dict) -> None:
    ws = wb.active
    ws.title = "Quick reference"
    style_title(ws, f"{verb.upper()} • complete modern conjugation", 7)

    ws.merge_cells("A2:G2")
    ws["A2"] = "Spanish verb study workbook • European Spanish forms included"
    ws["A2"].font = Font(italic=True, color="44546A")
    ws["A2"].alignment = Alignment(horizontal="center")

    key_forms = [
        ("Infinitive", nonfinite(data, "infinitivo", "infinitivo")),
        ("Gerund", nonfinite(data, "gerundio", "gerundio")),
        ("Past participle", nonfinite(data, "participo", "participo")),
    ]

    ws["A4"] = "Key forms"
    ws["A4"].font = Font(bold=True, color=WHITE)
    fill(ws["A4"], BLUE)
    ws.merge_cells("A4:B4")

    for row, (label, value) in enumerate(key_forms, start=5):
        ws.cell(row, 1, label).font = Font(bold=True, color="666666")
        ws.cell(row, 2, value).font = Font(bold=True, color=NAVY)

    table_row = 10
    headers = ["Person"] + [label for _, _, label in CORE_TENSES]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(table_row, col, header)
        cell.font = Font(bold=True, color=NAVY)
        fill(cell, LIGHT)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    tense_forms = [forms_for(data, mood, tense) for mood, tense, _ in CORE_TENSES]

    for person_index, (person_label, _) in enumerate(DISPLAY_PERSONS):
        row = table_row + 1 + person_index
        ws.cell(row, 1, person_label).font = Font(bold=True)
        for tense_index, values in enumerate(tense_forms, start=2):
            ws.cell(row, tense_index, values[person_index])
        if person_index % 2 == 0:
            for col in range(1, len(headers) + 1):
                fill(ws.cell(row, col), PALE)

    widths = {"A": 27, "B": 19, "C": 25, "D": 22, "E": 20, "F": 22}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    ws.freeze_panes = "B11"


def build_full_conjugation(wb: Workbook, verb: str, data: dict) -> None:
    ws = wb.create_sheet("Full conjugation")
    style_title(ws, f"Full conjugation of {verb.upper()}", 7)
    row = 3

    for section_name, colour, tenses in FULL_SECTIONS:
        section_bar(ws, row, section_name, colour)
        row += 1

        for mood, tense, label in tenses:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
            cell = ws.cell(row, 1, label)
            cell.font = Font(bold=True)
            fill(cell, GREY)
            row += 1

            ws.cell(row, 1, "Person").font = Font(bold=True, color=NAVY)
            ws.cell(row, 2, "Spanish form").font = Font(bold=True, color=NAVY)
            fill(ws.cell(row, 1), LIGHT)
            fill(ws.cell(row, 2), LIGHT)

            for (person_label, _), form in zip(DISPLAY_PERSONS, forms_for(data, mood, tense)):
                row += 1
                ws.cell(row, 1, person_label)
                ws.cell(row, 2, form)

            row += 2
        row += 1

    section_bar(ws, row, "IMPERATIVE • COMMANDS", ORANGE)
    row += 1
    for col, header in enumerate(("Person", "Affirmative", "Negative"), start=1):
        ws.cell(row, col, header).font = Font(bold=True, color=NAVY)
        fill(ws.cell(row, col), LIGHT)

    affirmative = forms_for(data, "imperativo", "afirmativo")
    negative = forms_for(data, "imperativo", "negativo")
    for index in range(1, len(DISPLAY_PERSONS)):
        row += 1
        ws.cell(row, 1, DISPLAY_PERSONS[index][0])
        ws.cell(row, 2, affirmative[index])
        ws.cell(row, 3, negative[index])

    row += 3
    section_bar(ws, row, "NON-FINITE FORMS", TEAL)
    row += 1
    for col, header in enumerate(("Form", "Spanish"), start=1):
        ws.cell(row, col, header).font = Font(bold=True, color=NAVY)
        fill(ws.cell(row, col), LIGHT)

    participle = nonfinite(data, "participo", "participo")
    nonfinite_rows = [
        ("Infinitive", nonfinite(data, "infinitivo", "infinitivo")),
        ("Compound infinitive", "haber " + participle),
        ("Gerund", nonfinite(data, "gerundio", "gerundio")),
        ("Compound gerund", "habiendo " + participle),
        ("Past participle", participle),
    ]
    for label, form in nonfinite_rows:
        row += 1
        ws.cell(row, 1, label)
        ws.cell(row, 2, form)

    ws.column_dimensions["A"].width = 33
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 30
    ws.freeze_panes = "A4"


def build_all_forms(wb: Workbook, verb: str, data: dict) -> None:
    ws = wb.create_sheet("All forms")
    style_title(ws, f"All forms returned for {verb.upper()}", 5)

    headers = ["Mood", "Tense", "Pronoun", "Form", "Alternatives"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(3, col, header)
        cell.font = Font(bold=True, color=WHITE)
        fill(cell, BLUE)

    row = 3
    for mood, tenses in data["moods"].items():
        for tense, entries in tenses.items():
            for entry in entries:
                forms = entry.get("c") or []
                row += 1
                ws.cell(row, 1, mood)
                ws.cell(row, 2, tense)
                ws.cell(row, 3, entry.get("pr", ""))
                ws.cell(row, 4, strip_pronoun(forms[0]) if forms else "")
                ws.cell(row, 5, " | ".join(forms[1:]))

    ws.auto_filter.ref = f"A3:E{row}"
    ws.freeze_panes = "A4"
    widths = {"A": 20, "B": 34, "C": 18, "D": 30, "E": 38}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def build_sources(wb: Workbook) -> None:
    ws = wb.create_sheet("Sources")
    ws.append(["Source", "URL"])
    ws.append(["Verbecc HTTP API", "https://verbe.cc/verbecc/conjugate/es/"])
    ws.append(["RAE Diccionario de la lengua española", "https://dle.rae.es/"])
    ws.sheet_state = "hidden"


def build_workbook(verb: str, payload: dict) -> bytes:
    data = normalise(payload)
    wb = Workbook()

    build_quick_reference(wb, verb, data)
    build_full_conjugation(wb, verb, data)
    build_all_forms(wb, verb, data)
    build_sources(wb)

    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    horizontal=cell.alignment.horizontal or "left",
                    vertical="center",
                    wrap_text=True,
                )

    wb["Quick reference"]["A1"].comment = Comment(
        "Conjugations retrieved from the Verbecc HTTP API. Review unusual or predicted verbs against the RAE.",
        "Verbo Magico",
    )

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
